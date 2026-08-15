from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\GCRA\nurse\GCON_2027_system_details_for_developers.docx")
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "5F6B76"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
RED = "9B1C1C"
GOLD = "7A5A00"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="D7DEE6", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)


def mark_header_row(table):
    if not table.rows:
        return
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def set_run_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def format_paragraph(paragraph, before=0, after=6, line=1.25, alignment=None, keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if alignment is not None:
        paragraph.alignment = alignment
    if keep_next:
        fmt.keep_with_next = True


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)
    return run


def add_body(doc, text, after=6):
    p = doc.add_paragraph(style="Normal")
    format_paragraph(p, after=after)
    add_text(p, text)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    format_paragraph(p, after=4)
    add_text(p, text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    format_paragraph(p, after=4)
    add_text(p, text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    add_text(p, text, size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_label_value(doc, label, value):
    p = doc.add_paragraph(style="Normal")
    format_paragraph(p, after=3, line=1.15)
    add_text(p, f"{label}: ", size=10, color=MUTED, bold=True)
    add_text(p, value, size=10, color=NAVY)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    header = table.rows[0].cells
    for index, text in enumerate(headers):
        set_cell_shading(header[index], LIGHT_BLUE)
        p = header[index].paragraphs[0]
        format_paragraph(p, after=0, line=1.1)
        add_text(p, text, size=9.5, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            p = cells[index].paragraphs[0]
            format_paragraph(p, after=0, line=1.12)
            add_text(p, str(value), size=9.2, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, label, text, fill=CALLOUT, label_color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent_dxa=120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    format_paragraph(p, after=2, line=1.2)
    add_text(p, f"{label}: ", size=10.5, color=label_color, bold=True)
    add_text(p, text, size=10.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color, before, after in ((1, 16, BLUE, 18, 10), (2, 13, BLUE, 14, 7), (3, 12, DARK_BLUE, 10, 5)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "GCON 2027 | Developer system handoff"
    format_paragraph(hp, after=0, line=1.0)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        set_run_font(run, size=9, color=MUTED, bold=True)
    footer = section.footer
    fp = footer.paragraphs[0]
    format_paragraph(fp, after=0, line=1.0)
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(fp, "Internal development reference | Page ", size=9, color=MUTED)
    add_page_field(fp)


def build():
    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph()
    format_paragraph(p, before=0, after=4, line=1.0)
    add_text(p, "DEVELOPER HANDOFF / SYSTEM DETAILS", size=10, color=GOLD, bold=True)
    p = doc.add_paragraph()
    format_paragraph(p, after=5, line=1.0)
    add_text(p, "GCON 2027 Nursing Application Platform", size=25, color=NAVY, bold=True)
    p = doc.add_paragraph()
    format_paragraph(p, after=16, line=1.15)
    add_text(p, "Current local MVP architecture, workflows, data model, API surface, annual intake configuration, and developer handoff", size=13, color=MUTED)

    meta = [
        ("Prepared for", "Gauteng Provincial Government / GCON development team"),
        ("Prepared from", "Current repository at C:\\GCRA\\nurse, the design specification, and the live presentation"),
        ("Document date", date.today().strftime("%d %B %Y").lstrip("0")),
        ("Current status", "Working React/Vite + Node local MVP plus a private ChatGPT Sites hosted demo"),
        ("Demo note", "Local demo role shortcuts are intentionally retained for the presentation; they are not the production identity solution"),
    ]
    table = doc.add_table(rows=0, cols=2)
    set_table_geometry(table, [1800, 7560])
    for label, value in meta:
        cells = table.add_row().cells
        set_cell_shading(cells[0], LIGHT_GRAY)
        p = cells[0].paragraphs[0]
        format_paragraph(p, after=0, line=1.1)
        add_text(p, label, size=9.5, color=NAVY, bold=True)
        p = cells[1].paragraphs[0]
        format_paragraph(p, after=0, line=1.1)
        add_text(p, value, size=9.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_callout(doc, "Executive summary", "The platform is a role-based nursing intake and placement prototype. It supports a learner qualification check and application journey, staff review and placement operations, and an organisation-scoped employer view. The local implementation has a real API boundary, workflow state transitions, encrypted local persistence, document upload checks, audit events, pathway-scoped academic ranking, and demo access. A private ChatGPT Sites deployment now provides a hosted presentation copy with the demo role pickers and seeded demo state. Production deployment still requires a managed database, private object storage, managed identity, malware scanning, external notification delivery, operational monitoring, and formal policy approval. No cross-pathway academic equivalency matrix is currently approved.")

    add_heading(doc, "1. What has been built", 1)
    add_body(doc, "The current build is a local MVP intended for stakeholder demonstration and workflow acceptance. The frontend is no longer dependent on mockStore.js for its active path; it calls the local API for state-changing operations. The local API uses an encrypted development adapter, so it behaves like a persisted application without claiming to be the production Azure/PostgreSQL implementation.")
    add_table(doc, ["Area", "Current capability", "Current boundary"], [
        ("Learner portal", "Landing page with the current advert and top navigation, server-rechecked qualification checker, pathway forms, ID-number access management, profile, remaining subjects, preferences, address confirmation, evidence upload, training and work history, review, receipt, correction request and final outcome views.", "Local demo profile values and local draft persistence remain. Production identity and policy configuration are still required."),
        ("Staff workspace", "Dashboard, submitted applications, academic ranks within each pathway, review queue, per-document verify/reject/request-correction actions, individual and mass decline decisions, linked in-app outcomes, shortlist, psychometric invitation and outcome, communications, capacity-controlled placements, withdrawals, termination-letter recording, reports and audit.", "The API is a Node local service rather than the planned NestJS/PostgreSQL production stack. Cross-pathway ranking remains unavailable without an approved equivalency matrix. Some seeded demo records have already been used during QA."),
        ("Employer portal", "Organisation-scoped dashboard, released candidates, placement offer responses, placement state visibility, approved campus directory, print/PDF view and campus-grouped Excel export.", "Employer access is limited to released/assigned GCON records. Future hospitals and clinics can be added through the organisation/membership model."),
        ("Access control", "Real email/password local sessions, learner account creation with an ID-number username, expiry handling, invitations, membership-derived roles, logout, plus retained local demo shortcuts for the presentation.", "Local passwords are development seeds. Production must use Microsoft Entra External ID or an equivalent managed identity provider."),
        ("Documents", "Upload intent, PDF/JPG/PNG metadata restrictions, size limits, checksum validation, private local document bytes, upload completion state and staff review metadata.", "The current provider is local encrypted storage. Production needs private Blob Storage, malware scanning and short-lived download links."),
        ("Notifications", "Communication batches and learner notifications are persisted with status/provider metadata and shown in staff communication screens.", "The local provider is a development adapter. Production needs an email/SMS provider, delivery callbacks and retry/dead-letter handling."),
    ], [1800, 3600, 3960])

    add_heading(doc, "2. Product scope and user journeys", 1)
    add_heading(doc, "Learner journey", 2)
    for item in [
        "Open the public GCON intake landing page and see the current annual intake name, open/close message and Gauteng-only messaging.",
        "Run the qualification checker before full registration. Supported pathways are Senior Certificate / Standard 10, National Senior Certificate / Grade 12, and NC(V) Level 4. Original pathway values are preserved and are not converted into a common score.",
        "If the learner qualifies, create an ID-number account through Access management, then continue to the pathway-specific profile and application form. If not, show the selected certificate and masked ID for reference while retaining only the minimum contact request.",
        "Capture personal details, South African ID, structured Gauteng address, school, result year, pathway-specific remaining subjects, preferences, previous nursing training, detailed work experience, references and the two approved evidence documents.",
        "Upload real PDF, JPG or PNG files through an API upload intent, direct content upload and completion/checksum step. The returned document state is retained in the draft and the learner can retry failures.",
        "Review the immutable submission snapshot, submit once, receive a daily reference number and receipt, and later see only correction requests or final outcomes rather than a detailed progress tracker.",
    ]:
        add_number(doc, item)

    add_heading(doc, "Staff journey", 2)
    for item in [
        "Open the admissions dashboard for the active intake cycle.",
        "Search applications and open the review queue for records that are under review or have been resubmitted after correction.",
        "Use academic rank only within the applicant's qualification pathway. APS, M score and NC(V) percentages are different measures and are not placed into a common intake-wide order.",
        "Review original pathway values and each evidence document. Each document can be verified, rejected, or marked correction required with a reason.",
        "Record the controlled application decision: approve for shortlist, request correction, decline with a controlled reason, use mass decline for eligible under-review/correction-requested records, or undo where allowed by the current workflow.",
        "Invite the psychometric test and later record passed, failed, no-show or pending outcome with optional score and notes.",
        "Prepare placement offers, record withdrawals, persist a termination-letter record, issue controlled communications, export reports, and inspect the append-only audit trail.",
    ]:
        add_number(doc, item)

    add_heading(doc, "Employer journey", 2)
    add_body(doc, "The employer view is organisation-scoped. It can show pipeline totals and candidates explicitly released for the GCON organisation. It can respond to a placement offer with an accepted or declined response and campus detail. It cannot perform staff eligibility decisions, browse unreleased candidates, see other organisations, or access unreleased documents.")

    add_heading(doc, "3. Architecture and runtime", 1)
    add_callout(doc, "Local runtime flow", "Browser (React/Vite) -> /v1 API on port 4000 -> domain workflow functions -> encrypted local state file and local private document bytes. Vite serves the frontend on port 4173 and proxies API calls during local development.")
    add_heading(doc, "ChatGPT Sites hosted demo", 2)
    add_body(doc, "The current frontend demo is published at https://gcon-nursing-intake.alistairljohanson.chatgpt.site. The Sites deployment uses the compiled Vite assets plus dist/server/index.js, a Cloudflare-compatible worker generated by tools/build-sites-worker.mjs. The hosted worker serves the learner, admissions and employer presentation surfaces and provides a deliberately limited demo API with seeded state so the retained role pickers continue to work outside the local machine. Access management is labelled Sites demo registration and creates only an expiring in-memory learner session; it does not retain passwords, identifiers or personal contact details.")
    add_callout(doc, "Hosted demo boundary", "The Sites demo is not connected to the local encrypted API, PostgreSQL, Azure Blob Storage or local uploaded documents. Hosted state-changing actions are not a production persistence boundary. Use the local app for API, real-file upload, audit, database and document-storage acceptance. The public Sites demo registration is temporary and must not be used with real personal information.", fill="FFF8E8", label_color=GOLD)
    add_table(doc, ["Layer", "Current implementation", "Production target"], [
        ("Frontend", "React components in src/main.jsx plus CompleteApplicantViews, CompleteWorkspacePages and CompleteEmployerWorkspace; CSS in src/styles.css.", "Keep React, split into feature modules, add typed API client, robust error surfaces, accessibility and production build pipeline."),
        ("API", "Node http server in server.mjs with JSON routes, request IDs, predictable error envelope, role checks and API contract in openapi/v1.json.", "Typed NestJS API with generated/validated OpenAPI, middleware for identity, idempotency, transactions and observability."),
        ("Domain", "server/domain.mjs contains statuses, permissions, qualification-independent workflow actions, audit writes and organisation scoping.", "Move domain logic into service modules with PostgreSQL transactions and explicit state-transition policies."),
        ("Authentication", "server/auth.mjs provides local scrypt password verification, expiring sessions, invitations and membership-derived roles. Demo tokens remain available only in demo mode.", "Microsoft Entra External ID or equivalent managed identity provider, secure cookies or validated JWTs, MFA for staff, account recovery and conditional access."),
        ("Storage", "server/storage.mjs encrypts the local state file and stores document bytes in a local documents directory.", "Azure PostgreSQL Flexible Server for system of record; private Azure Blob Storage for documents; Key Vault for secrets."),
        ("Operations", "Health endpoints and provider status exist for local verification. Local API and frontend logs are available during development.", "Application Insights/Log Analytics, alerts, backups, restore rehearsals, incident response and privacy-safe telemetry."),
        ("Hosted demo", "ChatGPT Sites deployment at gcon-nursing-intake.alistairljohanson.chatgpt.site. npm run build generates the static/demo worker and .openai/hosting.json stores only the Sites project ID. Sites demo registration creates a temporary in-memory session and retains no password or personal contact details.", "Deploy the approved production frontend against the production API, managed identity, database, private document storage and external notification services."),
    ], [1800, 3600, 3960])

    add_heading(doc, "4. Roles, memberships and permission boundaries", 1)
    add_table(doc, ["Role", "Can do", "Cannot do"], [
        ("Learner", "Own application draft, own uploads, own receipt, correction response and final outcome.", "Read other applications, staff audit data, employer data or unreleased documents."),
        ("Staff reviewer", "Review applications and documents in the GCON organisation, record decisions, shortlist, invite and record psychometric outcomes, manage placements and communications.", "Act outside organisation scope or release records to another organisation."),
        ("Staff supervisor", "Reviewer permissions plus membership invitations and supervisory administration.", "Bypass workflow or access another organisation without an explicit membership."),
        ("Employer member/coordinator", "See only released/assigned GCON candidates, placement work and responses.", "Make eligibility decisions, see unreleased documents or browse other organisations."),
        ("Platform administrator", "Platform-wide administrative operations subject to audit.", "Operate outside audit or bypass retention/privacy policy."),
        ("Demo shortcut", "Open the learner, staff or employer demo session for tomorrow's presentation.", "It is not a production authentication mechanism and must be disabled or removed from production configuration."),
    ], [1800, 3900, 3660])

    add_heading(doc, "5. Qualification and data rules", 1)
    add_body(doc, "The qualification checker preserves the learner's original reported values. It does not silently convert percentages to APS, create certificate equivalencies, or turn an indicative result into the official staff decision. The current rules in src/qualification.js are the working prototype rules and must be approved before production use.")
    add_callout(doc, "Academic ranking boundary", "Academic rank is calculated separately within the NSC / Grade 12, Senior Certificate and NC(V) Level 4 pathways. An APS, M score and NC(V) percentage are not numerically comparable, so the MVP does not produce one cross-pathway order. Cross-pathway ranking must remain unavailable until GCON approves a versioned equivalency matrix covering formulas, included subjects, weights, maxima, rounding, ties, upgraded or missing results, validation evidence, effective cycle, approvers and appeals.", fill="FFF8E8", label_color=GOLD)
    add_table(doc, ["Pathway", "Current checker rule"], [
        ("NSC / Grade 12", "English Level 4 or higher; Life Sciences Level 4 or higher; Mathematics Level 4 or higher OR Maths Literacy Level 5 or higher; reported APS 27 or higher."),
        ("Senior Certificate", "English, Biology and Mathematics in the supported HG/SG pass bands; M score 17 or higher."),
        ("NC(V) Level 4", "Fundamental subjects at 50% or higher; SA Health Care System, Public Health, The Human Body and Mind, and Community Oriented Primary Care at 60% or higher."),
        ("Academic ranking", "Eligible records are ordered only against records in the same qualification pathway, using the pathway's original reported academic value. No common score or cross-pathway rank is approved."),
        ("Data minimisation", "Non-qualifying contact capture retains minimal contact details and does not retain pathway results, APS or other qualification evidence."),
        ("Identity protection", "Learner-facing application responses mask the South African ID. Staff review requires controlled access to the original protected data."),
    ], [2200, 7160])

    add_heading(doc, "6. Core records and workflow states", 1)
    add_table(doc, ["Record", "Important fields / purpose"], [
        ("Users and identities", "User ID, name, email, local development credential or external identity reference, active roles and organisation IDs."),
        ("Organisations and memberships", "Organisation, active membership, role, scope and future hospital/clinic expansion point."),
        ("Intake cycle", "Cycle ID, name/season, open date, close date, advert status, reference date, policy status, document labels and any future approved equivalency-matrix version."),
        ("Application", "Reference, owner, pathway, profile, remaining subjects, preferences, original pathway values, training history, work experience, status, submittedAt, correction request, outcome and placement information."),
        ("Submission snapshot", "The submitted application state becomes the learner's receipt-backed record. Corrections are separate controlled changes and resubmission events."),
        ("Document", "Document ID, application reference, type/label, filename, content type, size, checksum, scan state, review state, reviewer and object key."),
        ("Review task / decision", "Queue assignment, staff decision, controlled reason, correction request, shortlist release and audit event."),
        ("Psychometric record", "Invitation status, scheduled time, outcome, score, notes, recorded by and recorded time."),
        ("Placement", "Campus, placement-ready state, employer response, final placed state, withdrawal and termination-letter record."),
        ("Communication / notification", "Audience, template, recipient count, provider, delivery status, queued/sent times, recipient application and learner visibility."),
        ("Audit event", "Append-only event, timestamp, actor, actor user ID, organisation, reference, document ID and reason/details."),
    ], [2400, 6960])

    add_heading(doc, "7. API surface", 1)
    add_body(doc, "All protected routes require identity and validate the active role, organisation scope, record ownership or release, and the allowed workflow transition before writing state or audit data. State-changing routes return an error envelope containing a request ID. Submission uses an idempotency key and duplicate application guards.")
    add_table(doc, ["Area", "Routes currently implemented"], [
        ("Health / contract", "GET /v1/health; GET /v1/openapi.json; GET /v1/storage/status; GET /v1/state"),
        ("Authentication", "POST /v1/auth/login; POST /v1/auth/logout; GET /v1/auth/session; POST /v1/auth/invitations; GET /v1/auth/invitations; POST /v1/auth/invitations/accept; POST /v1/auth/demo-login; POST /v1/auth/learner/register"),
        ("Qualification", "POST /v1/qualification/evaluate; returns the selected pathway's original score/result without a cross-pathway conversion."),
        ("Intake", "GET/PATCH /v1/intake/current"),
        ("Learner application", "GET/PATCH /v1/applications/me; POST /v1/applications/submit; POST /v1/applications/non-qualifier; GET/POST /v1/applicant/chat"),
        ("Documents", "POST /v1/documents/upload-intent; PUT /v1/documents/{id}/content; POST /v1/documents/{id}/complete; POST /v1/documents/{id}/review"),
        ("Staff review", "GET /v1/reviews/queue; POST /v1/reviews/{ref}/decision; POST /v1/reviews/mass-decline; POST /v1/shortlists"),
        ("Psychometric", "POST /v1/interviews; PATCH /v1/interviews/{ref}/outcome"),
        ("Placement", "POST /v1/placements/{ref}/prepare; POST /v1/placements/{ref}/response; PATCH /v1/applications/{ref}/withdrawal; POST /v1/placements/{ref}/termination-letter"),
        ("Employer", "GET /v1/employer/dashboard; GET /v1/employer/candidates"),
        ("Communication / audit", "POST /v1/communications/issue; GET /v1/communications; GET /v1/notifications; GET /v1/audit"),
    ], [2200, 7160])

    add_heading(doc, "8. Annual intake and configurable requirements", 1)
    add_callout(doc, "Required product capability", "GCON runs annual intakes. Staff administrators must be able to create or activate a new intake season, change the basic dates and advert status, and update approved requirement labels and rules without editing source code. Every application and submission snapshot must retain the cycle version it was created under.", fill="FFF8E8", label_color=GOLD)
    add_heading(doc, "What can change in the intake configuration", 2)
    for item in [
        "Cycle identity: year, season/name, public advert title and cycle status.",
        "Availability: opening date/time, closing date/time, timezone, late-application policy and reference-number date basis.",
        "Eligibility wording: pathway names, approved subject labels, thresholds, APS/M score rules and NC(V) percentage rules, with an approval state and effective version. Any cross-pathway equivalency matrix must be a separate approved, versioned policy object rather than an inferred conversion.",
        "Evidence requirements: document type labels, required/optional flag, allowed MIME types, maximum size, review guidance and retention class.",
        "Application content: address/Gauteng confirmation wording, preference list, references, training and experience questions, and support contact details.",
        "Communications: approved templates, sender identity, delivery channels, audience definitions and outcome language.",
        "Operational settings: review queue rules, placement campuses, organisation release rules, retention schedule and reporting dimensions.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Current state and implementation recommendation", 2)
    add_body(doc, "The local API already supports the active cycle name, open date, close date and Published/Paused advert status through PATCH /v1/intake/current. The current checker rules, document labels and within-pathway ranking behavior are still code-defined and provisional. No cross-pathway equivalency is configured. For production, introduce versioned IntakeCycle and IntakeRequirement records, an administrator-only configuration screen, approval/sign-off before publishing, and immutable linkage from every application to the configuration version used at submission.")
    add_table(doc, ["Change type", "Local MVP today", "Production acceptance requirement"], [
        ("Open/close dates", "PATCH /v1/intake/current; banner and intake-open validation use cycle dates.", "Timezone-aware date/time, audit history, future-cycle scheduling and close-date enforcement in database transactions."),
        ("Advert status", "Published or Paused; protected staff route.", "Role-restricted publish workflow, preview, approval and effective timestamp."),
        ("Requirements", "Rules in src/qualification.js and document labels in the cycle seed.", "Versioned, approved, data-driven rules with validation and a no-retroactive-change policy."),
        ("Cross-pathway equivalency", "Not configured; the MVP ranks only within each pathway.", "Approved matrix with formulas, weights, maxima, rounding, tie rules, validation, effective cycle, approvers and immutable score snapshots."),
        ("New annual intake", "Requires local data/configuration preparation.", "Create next cycle without overwriting prior cycles; preserve prior submissions, references, reports and audit history."),
    ], [2000, 3300, 4060])

    add_heading(doc, "9. Document and notification processing", 1)
    add_heading(doc, "Document upload sequence", 2)
    for item in [
        "Learner requests an upload intent with document type, filename, content type and size.",
        "API validates PDF, JPEG or PNG type and a maximum size of 10 MB, creates a document record and returns a short-lived local upload route.",
        "Browser uploads the file bytes. The API validates the declared size, computes SHA-256 and marks the document pending scan/review.",
        "Learner completes the upload with the checksum. The document is available to staff review; production will insert malware scanning before a clean state.",
        "Staff records one of verified, rejected or correction_required, with reason, reviewer and time. The application evidence and audit event are updated together.",
    ]:
        add_number(doc, item)
    add_heading(doc, "Notification sequence", 2)
    add_body(doc, "The local communication action creates a persisted batch, creates per-recipient learner notifications, records provider and delivery status metadata, and exposes the batch in the staff communication centre. The provider is explicitly named local-development-adapter so it is not mistaken for external delivery. Production must add an email/SMS provider, delivery webhooks, retries, failure reasons, status reconciliation, PII-safe logs and operational alerts.")

    add_heading(doc, "10. Security, privacy and audit posture", 1)
    for item in [
        "Organisation scoping is applied to staff and employer reads and writes. Employers see only released or assigned candidates.",
        "Learner responses mask identifiers and do not expose internal organisation IDs, release lists or staff audit data.",
        "Documents are private by default; the production design must use private Blob containers and short-lived links rather than public URLs.",
        "Passwords in the local adapter use scrypt hashing and sessions expire. Production credentials must not use seeded passwords.",
        "Audit events are append-only in the current domain model and accompany workflow actions such as submission, review, correction, shortlist, placement, withdrawal, communications and termination letters.",
        "POPIA-aligned safeguards, retention, legal basis, data subject rights and breach response require formal approval from the responsible department.",
    ]:
        add_bullet(doc, item)
    add_callout(doc, "Do not promote the local adapter", "The encrypted local store is suitable for demo and development only. It is not a substitute for PostgreSQL, Blob Storage, managed identity, Key Vault, malware scanning, backups, monitoring or an approved retention schedule.", fill="FDECEC", label_color=RED)

    add_heading(doc, "11. QA and demo readiness", 1)
    add_table(doc, ["Verification area", "Current evidence / next check"], [
        ("Automated tests", "40 Node tests currently pass across domain, auth, qualification, upload-policy and employer-export suites. Extend with document magic-byte fixtures, cross-organisation denial, equivalency-policy rejection and notification status failure/retry cases."),
        ("Build", "npm run build passes, compiles the Vite frontend and generates the Sites worker at dist/server/index.js. Re-run after any demo UI change."),
        ("Browser QA", "Use the running local frontend at http://127.0.0.1:4173/ and API health at http://127.0.0.1:4000/v1/health. Re-run learner upload, staff document actions, employer placement response and session expiry flows before the presentation."),
        ("Hosted demo QA", "Open https://gcon-nursing-intake.alistairljohanson.chatgpt.site, choose Applicant, Admissions or Employer demo, and confirm the corresponding view loads. Hosted data is presentation-only and must not be used to validate persistence."),
        ("Mobile QA", "Run learner submission and staff review at a narrow mobile viewport, including upload controls, queue cards, tables and modal login."),
        ("Permission QA", "Verify learner self-only access, employer release-only access, cross-organisation denial, staff-only actions and expired-session 401 handling."),
        ("Duplicate submission", "Submit once with an idempotency key, retry the same request, then attempt a new submission and confirm the API prevents a second submitted application."),
        ("Real file QA", "Use one valid PDF, one valid JPG and one valid PNG and verify intent, byte upload, checksum, completion and staff review state for each."),
    ], [2200, 7160])
    add_heading(doc, "Presentation demo accounts", 2)
    add_table(doc, ["Demo role", "Local access"], [
        ("Applicant", "Use the Applicant demo shortcut, or email lerato.mokoena@email.com with password learner-demo."),
        ("Admissions", "Use the Admissions demo shortcut, or email thandi.mokoena@gcon.example with password staff-demo."),
        ("Employer", "Use the Employer demo shortcut, or email placements@gcon.example with password employer-demo."),
        ("Admin", "Local account exists for API/admin testing: admin@gcon.example with password admin-demo."),
        ("ChatGPT Sites", "Private hosted demo URL: https://gcon-nursing-intake.alistairljohanson.chatgpt.site. Sign in with the owning ChatGPT account, then use the retained demo role pickers."),
    ], [2200, 7160])
    add_body(doc, "The demo shortcuts intentionally remain in the local login dialog and in the private hosted presentation copy. Local shortcuts use the local demo endpoint; the hosted copy uses its limited seeded demo worker. Both are presentation-only and must be disabled by setting GCON_AUTH_MODE to a non-demo production mode and configuring the external identity adapter for production.")

    add_heading(doc, "12. Production handoff and remaining work", 1)
    for item in [
        "Confirm the official closing date, post-submission edit policy, exact document labels, employer permissions, retention schedule, support owner, sender identity, Azure subscription and region, and approved qualification wording. Keep cross-pathway ranking disabled unless GCON also approves a complete equivalency matrix.",
        "Replace the Node local adapter with a typed NestJS service layer, PostgreSQL migrations, transactional workflow actions and database-backed organisation scope.",
        "Bind authentication to Microsoft Entra External ID or an equivalent provider, add staff MFA, recovery, invitation lifecycle and role/membership administration.",
        "Bind private Azure Blob Storage, upload intents, checksum validation, malware scanning, retention and short-lived download links.",
        "Move annual intake requirements into versioned, approved configuration and deliver an administrator configuration screen.",
        "Complete external notification delivery, status callbacks, retries, failure handling and PII-free observability.",
        "Finish mobile and browser end-to-end coverage, load/concurrency tests for 50,000-100,000 records, backup/restore rehearsals and an intake pilot.",
        "After API/database acceptance, remove demo-only data paths and prototype-only components from the production build. Keep the demo shortcuts only in the isolated presentation/local configuration until the demo is complete; do not treat the hosted Sites worker as the production API.",
    ]:
        add_number(doc, item)

    add_heading(doc, "13. Developer file map and commands", 1)
    add_table(doc, ["File / command", "Purpose"], [
        ("src/main.jsx", "App shell, learner portal, staff navigation, login/demo dialog and presentation composition."),
        ("src/api.js", "Frontend API client, session persistence, upload sequence and workflow action calls."),
        ("src/qualification.js", "Current qualification pathways and prototype rule evaluation."),
        ("src/CompleteApplicantViews.jsx", "Learner review, receipt, correction and outcome presentation."),
        ("src/CompleteWorkspacePages.jsx", "Staff applications, review queue, document actions, shortlist, communications, reports, audit and placements."),
        ("src/CompleteEmployerWorkspace.jsx", "Organisation-scoped employer dashboard and placement response view."),
        ("server.mjs", "Local HTTP API, routing, auth boundary, upload handling and persistence calls."),
        ("server/domain.mjs", "Domain records, permissions, state transitions, notifications and audit events."),
        ("server/auth.mjs", "Local session, invitation, membership and demo auth adapter."),
        ("server/storage.mjs", "Encrypted local state and document storage adapter."),
        ("openapi/v1.json", "Machine-readable API contract."),
        ("npm test", "Run Node domain, auth, qualification, upload-policy and employer-export tests."),
        ("npm run build", "Build the Vite frontend for a production-style compile check."),
        ("tools/build-sites-worker.mjs", "Generate the Cloudflare-compatible static/demo worker used by the private ChatGPT Sites presentation deployment."),
        (".openai/hosting.json", "Sites project binding. Contains only the opaque project ID; credentials are not stored in the repository."),
        ("npm run dev:api", "Run the local API on port 4000."),
        ("npm run dev -- --host 127.0.0.1 --port 4173", "Run the local Vite frontend."),
    ], [3100, 6260])

    add_heading(doc, "Source and interpretation notes", 1)
    add_body(doc, "This document describes the state of the local repository and its local API, not a claim that production infrastructure has already been deployed. The design specification is the primary product source. The older presentation note about no tracking is interpreted as minimal learner status: receipt, correction requests and final outcome, without a detailed progress tracker. Qualification wording, document labels, retention, legal compliance and any cross-pathway equivalency matrix remain subject to departmental approval.")
    add_label_value(doc, "Primary specification", r"C:\Users\22524231\OneDrive - Gauteng Provincial Government\Desktop\GCON_2027_application_platform_design_specification.pdf")
    add_label_value(doc, "Presentation reference", r"C:\Users\22524231\OneDrive - Gauteng Provincial Government\Desktop\GCON_2027_live_presentation.pptx")
    add_label_value(doc, "Repository", r"C:\GCRA\nurse")
    add_label_value(doc, "Hosted demo", "https://gcon-nursing-intake.alistairljohanson.chatgpt.site (private ChatGPT Sites deployment)")

    for table in doc.tables:
        mark_header_row(table)

    doc.core_properties.title = "GCON 2027 Nursing Application Platform - System Details"
    doc.core_properties.subject = "Developer handoff and annual intake configuration reference"
    doc.core_properties.author = "GCON development handoff"
    doc.core_properties.comments = "Prepared from the local MVP implementation and supplied product references."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
