from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "GCON_2027_Chatbot_FAQ_Training.docx"
BLUE = RGBColor(11, 45, 82)
MID_BLUE = RGBColor(46, 116, 181)
PALE_BLUE = "EAF2F7"
LIGHT_BLUE = "E8EEF5"
GREY = RGBColor(92, 103, 112)


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_run(paragraph, text, bold=False, color=None, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return run


def add_bullet(document, text, level=0):
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(text)
    return paragraph


def add_number(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(text)
    return paragraph


def add_faq(document, question, answer, bullets=None, note=None):
    document.add_heading(question, level=3)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.add_run(answer)
    if bullets:
        for bullet in bullets:
            add_bullet(document, bullet)
    if note:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.space_after = Pt(8)
        add_run(paragraph, "Operator note: ", bold=True, color=GREY)
        add_run(paragraph, note, color=GREY)


def add_callout(document, title, text, fill=PALE_BLUE):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(4)
    add_run(paragraph, title, bold=True, color=BLUE, size=11)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    add_run(paragraph, text, size=10)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(34, 34, 34)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in (
        ("Title", 25, BLUE, 0, 8),
        ("Heading 1", 16, MID_BLUE, 18, 8),
        ("Heading 2", 13, MID_BLUE, 14, 6),
        ("Heading 3", 11.5, BLUE, 10, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15


def configure_section(section):
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(0)
    add_run(paragraph, "GAUTENG COLLEGE OF NURSING  |  2027 INTAKE", bold=True, color=BLUE, size=8.5)

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    add_run(paragraph, "Chatbot training reference  •  Internal operational use  •  Reviewed 31 August 2026", color=GREY, size=8)


def add_metadata_table(document):
    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    values = [
        ("Audience", "Chatbot trainers, admissions staff, service desk and employer-support staff"),
        ("Scope", "GCON 2027 nursing intake application and placement workflows"),
        ("Status", "Training reference; operational facts should be reviewed when intake policy changes"),
        ("Owner", "GCON admissions and application operations"),
    ]
    for row, (label, value) in zip(table.rows, values):
        shade(row.cells[0], LIGHT_BLUE)
        set_cell_margins(row.cells[0])
        set_cell_margins(row.cells[1])
        add_run(row.cells[0].paragraphs[0], label, bold=True, color=BLUE, size=9.5)
        add_run(row.cells[1].paragraphs[0], value, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    document = Document()
    configure_styles(document)
    for section in document.sections:
        configure_section(section)

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(12)
    title.add_run("GCON 2027 Chatbot Training FAQ")
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    add_run(subtitle, "Approved question-and-answer reference for applicant, admissions and employer support", color=GREY, size=12)
    add_metadata_table(document)

    add_callout(
        document,
        "Important environment boundary",
        "The public portal is a stakeholder demo. Do not put real identity numbers, certificates, contact details or passwords into the hosted demo. The local deployment is the operational path for encrypted application state, document storage and authenticated API workflows.",
    )

    document.add_heading("How the chatbot should answer", level=1)
    document.add_paragraph("Use this document as the approved source for common questions. Keep answers short, factual and action-oriented. When the answer depends on a person’s record, staff decision or current intake setting, direct the user to the portal or an authorised GCON support channel.")
    for item in [
        "Do not make or imply a final admissions decision. The qualification checker is an eligibility aid; staff verify evidence and make workflow decisions.",
        "Do not expose application records, review notes, identity numbers, certificates or placement information to another person.",
        "Do not ask users to send passwords or identity documents in chatbot messages. Tell them where to use the secure portal instead.",
        "Use the current intake name and dates shown in the portal. If the intake is paused or the dates are unclear, escalate to admissions staff.",
    ]:
        add_bullet(document, item)

    document.add_heading("1. General and access", level=1)
    add_faq(document, "What is the GCON nursing intake portal?", "It is the Gauteng College of Nursing application portal for the 2027 intake. Applicants use it to check pathway eligibility, create an account, complete a profile, upload evidence, review their application and submit it. Staff use the workspace to verify and process applications. Employers only see candidates released to them for placement work.")
    add_faq(document, "Who can use the portal?", "Applicants can create and submit their own application. Admissions staff can review applications and manage intake workflow. Employer users can respond to placement requests for candidates explicitly released to their organisation. Access is role-scoped.")
    add_faq(document, "What should I do if I cannot log in?", "Check that you are using the email address and password created for the portal. If the problem continues, use the authorised GCON support channel. Never share your password in the chatbot or by email.", note="For the local demo, the supplied demonstration accounts are for testing only and are not production credentials.")
    add_faq(document, "Is the public portal connected to the live admissions database?", "No. The public copy is a stakeholder demo. It is not the production source of truth for applications, encrypted local state, documents or PostgreSQL/Azure Blob storage. Use the approved local deployment for operational processing.")

    document.add_heading("2. Applicant journey", level=1)
    add_faq(document, "How do I apply?", "Complete the application in the following order:", bullets=[
        "Check your qualification pathway and minimum requirements.",
        "Create an account and sign in.",
        "Complete all required profile and contact fields.",
        "Add education, previous training, work experience and references where applicable.",
        "Upload both required documents.",
        "Open Review / submit application, check the summary, then select Submit application.",
    ])
    add_faq(document, "What happens if I close the browser before saving?", "Save each profile section before leaving the page. Applicants should not rely on closing and reopening the browser to preserve unsaved entries. If information appears missing, return to the relevant section and save it before continuing.")
    add_faq(document, "Can I submit more than one application?", "No. The system prevents multiple applications for the same identity number within the same intake. If you believe an earlier application was created incorrectly, contact authorised admissions support rather than creating another record.")
    add_faq(document, "How do I know whether my application was submitted?", "Open Review / submit application. A submitted application should show its submitted state and reference. If the page still asks you to submit, check the required fields and documents, then select Submit application.")
    add_faq(document, "What are the four campus choices?", "The current campus directory contains four choices:", bullets=[
        "Ann Latsky Campus",
        "Chris Hani Baragwanath Campus",
        "SG Lourens Campus",
        "Bonalesedi Campus",
    ])

    document.add_heading("3. Qualification criteria", level=1)
    add_faq(document, "What are the NSC / Grade 12 minimum requirements?", "The current checker requires English Level 4 or higher, Life Sciences Level 4 or higher, and either Mathematics Level 4 or higher or Mathematical Literacy Level 5 or higher. The minimum APS is 27. Staff still verify the uploaded certificate and supporting details before making a decision.")
    add_faq(document, "What are the Senior Certificate requirements?", "The Senior Certificate pathway uses approved Higher Grade and Standard Grade pass bands for English, Biology and Mathematics. The minimum M score is 17. The exact result must be checked against the pathway rules and verified by staff.")
    add_faq(document, "What are the NC(V) Level 4 requirements?", "The NC(V) Level 4 pathway requires Fundamentals at 50% or higher and each named vocational subject at 60% or higher. Staff verify the certificate and subject evidence.")
    add_faq(document, "What does the APS or M-score message mean?", "If the checker says the result is below the requirement, it means the entered or verified result does not meet the current pathway threshold. For NSC / Grade 12, the APS minimum is 27. For Senior Certificate, the M score minimum is 17. The message is not a final admissions decision.")
    add_faq(document, "Can I compare APS, M score and NC(V) percentages directly?", "No. These are pathway-specific measures. The system applies the relevant rules within the selected pathway and does not use one pathway’s numerical measure as a substitute for another.")

    document.add_heading("4. Profile, documents and submission", level=1)
    add_faq(document, "Which profile fields are required?", "Complete the required identity and contact information, including name, surname, email, mobile/cell number and address. Complete the education and supporting sections that apply to your pathway. The portal uses saved profile information in the application review summary.")
    add_faq(document, "Which documents must I upload?", "Upload both a Certified copy of ID and a Statement of results / certificate. The accepted formats are PDF, JPG and PNG. Wait for the upload to complete before leaving the page, then confirm that both documents appear in the review area.")
    add_faq(document, "Why is my work experience not shown in the review summary?", "Return to the work-experience section, check that each record has been saved, and reopen the review summary. A work record should include the employer, job title, dates, current-employer selection, responsibilities, reason for leaving where relevant and reference details. If a saved record is still missing, escalate to admissions support with the application reference.")
    add_faq(document, "What previous nursing training can I select?", "Choose one of the three specific nursing-training options shown in the portal. The current form does not use an open-ended Other option. If none of the options describes your training, contact admissions support before submitting.")
    add_faq(document, "How do I add another subject or complete the education section?", "Complete the current subject fields and use Submit / Add Remaining Subject when the form presents that action. Continue until all relevant subjects are recorded, then save the section and check the review summary.")
    add_faq(document, "What should I do if an intake or document label looks wrong?", "Do not guess or select an unrelated option. Record the application reference and contact authorised admissions support. Staff with the correct role can review cycle-scoped intake settings and approved evidence labels.")

    document.add_heading("5. Staff workflow", level=1)
    add_faq(document, "What does admissions staff do after an application is submitted?", "Staff open the Applications workspace, filter or search for the application, review the applicant profile and uploaded evidence, confirm the relevant pathway rules, and record the appropriate workflow decision. Review actions and decisions should remain traceable.")
    add_faq(document, "How should staff rank applicants?", "Rank applicants within the relevant pathway using the approved pathway-specific measures and verified evidence. Do not combine APS, M score and NC(V) percentages into one cross-pathway score.")
    add_faq(document, "What is the review queue?", "The review queue is the staff worklist for submitted applications needing verification or a decision. Use the application reference, filters and search to locate records, then open the full review summary before recording a decision.")
    add_faq(document, "Who can change intake settings?", "Cycle-scoped intake settings such as season name, dates, publication state, evidence labels and requirement wording should only be changed by an authorised supervisor or administrator. Changes should be checked against the audit trail.")
    add_faq(document, "What does Shortlisted mean?", "Shortlisted means staff have selected the applicant for the next placement-related step. It is not the same as a final placement or admission confirmation. The applicant’s next action should follow the current official communication.")
    add_faq(document, "What does Placement ready mean?", "Placement ready means the record has been prepared for placement handling and may be released to the appropriate employer workflow, subject to staff controls. It does not mean the employer has accepted the candidate.")

    document.add_heading("6. Employer workflow and reports", level=1)
    add_faq(document, "What can an employer user see?", "Employer users see only candidates explicitly released to their organisation for placement work. They can respond to placement requests and record the campus response. They do not see unreleased applicants, sensitive review evidence or staff eligibility decisions.")
    add_faq(document, "How do I respond to a placement request?", "Open the released candidate worklist, review the placement request, record the organisation’s response and submit it. The response should remain tied to the candidate and be traceable in the placement workflow.")
    add_faq(document, "How does the Excel export work?", "In Campus assignment reports, select Export Excel workbook (separate tabs). The workbook contains a Summary tab and a separate tab for each campus/status combination, including empty groups. It includes candidate and placement-contact fields for the released records in the current employer scope.")
    add_faq(document, "Which statuses have separate export tabs?", "The current employer workbook creates tabs for Shortlisted, Placement ready and Placed for each campus. It also includes the Summary tab with campus, status and applicant counts.")
    add_faq(document, "Why is a campus missing from the report?", "The current campus directory has four campuses. If a campus is not visible, refresh the page and confirm that the current intake settings and employer scope are loaded. If it is still missing, record the screen and escalate to the application administrator.")
    add_faq(document, "Can I print campus student lists?", "Use Print / save PDF from Campus assignment reports. Check the print preview before saving, especially the campus heading, status, contact details and page breaks. Do not distribute a list outside the approved organisation scope.")

    document.add_heading("7. Communications and support", level=1)
    add_faq(document, "Can messages include the applicant’s name?", "Yes. Approved communication templates can use the variables {{name}}, {{surname}} and {{id}}. The system resolves these values for the intended recipient. Staff should check the preview and recipient scope before sending.")
    add_faq(document, "What should a chatbot say when it cannot answer?", "Say: “I do not have enough information to answer that safely. Please contact authorised GCON admissions support and include your application reference, but do not send your password or identity documents in chat.”")
    add_faq(document, "What information should I include in a support request?", "Include the application reference, the page or action that failed, the approximate time, and a short description of what you expected to happen. Do not include a password, full identity number or copies of certificates in an unsecured message.")
    add_faq(document, "When must the chatbot escalate?", "Escalate when the user asks for a final eligibility or admissions decision, requests another person’s application information, reports a duplicate identity-number block that appears incorrect, cannot upload required documents, sees a missing saved section, or reports a payment/security/privacy concern.")

    document.add_heading("Escalation checklist", level=1)
    document.add_paragraph("Before handing a case to staff, capture only the minimum operational context needed to investigate:")
    for item in [
        "Applicant or employer role, without requesting sensitive credentials.",
        "Application reference, if available.",
        "Current page and action, for example Review / submit application or Export Excel workbook (separate tabs).",
        "Exact error text and whether the issue is repeatable.",
        "Browser/device context and approximate time.",
        "Whether the issue affects one record or multiple users.",
    ]:
        add_bullet(document, item)
    add_callout(document, "Never include in a chatbot ticket", "Passwords, one-time codes, full identity numbers, certificate images, private review notes or another applicant’s personal information.", fill="F6E9E7")

    document.add_heading("Source and maintenance notes", level=1)
    document.add_paragraph("This reference was compiled from the GCON nursing intake application’s current user interface, workflow code and operational README on 31 August 2026. Review it whenever intake policy, qualification thresholds, campus directory, roles, document requirements or message variables change.")
    for item in [
        "Applicant and staff workflow: README.md and the complete applicant/staff view modules.",
        "Qualification and campus rules: server/domain.mjs and application workflow components.",
        "Employer scope and exports: CompleteEmployerWorkspace.jsx and employerExports.js.",
        "Environment boundary: README.md and the configured local API/deployment settings.",
    ]:
        add_bullet(document, item)

    document.core_properties.title = "GCON 2027 Chatbot Training FAQ"
    document.core_properties.subject = "FAQ reference for GCON nursing intake chatbot training"
    document.core_properties.author = "Gauteng College of Nursing application operations"
    document.core_properties.keywords = "GCON, nursing intake, FAQ, chatbot, admissions, employer portal"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
